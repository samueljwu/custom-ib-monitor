from openpyxl import load_workbook

from app.storage.excel_io import export_workbook
from app.ui import COLUMNS




def test_export_keeps_formula_like_user_text_as_text(tmp_path):
    path = tmp_path / "safe.xlsx"
    export_workbook(path, [{"display_ticker": "=1+1"}], [], [], COLUMNS)
    book = load_workbook(path)
    try:
        headings = [cell.value for cell in book["Portfolio"][1]]
        cell = book["Portfolio"].cell(2, headings.index("Instrument") + 1)
        assert cell.value == "=1+1"
        assert cell.data_type == "s"
    finally:
        book.close()
